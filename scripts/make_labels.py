"""xlsx с разметкой -> data/labels_clean.csv (локально).

  python scripts/make_labels.py путь\\к\\open_labels.xlsx

Понимает и исходную таблицу организаторов (имя файла | таймстемп | нарушение / без нарушения),
и наш datathon_labels.xlsx (лист «разметка»).
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HOLDOUT = {"Х-2", "Х-4", "Х-8", "Х-10", "Х-15", "Х-17", "Х-23", "Х-32", "Х-39"}
PEOPLE = ["Чел 1", "Чел 2", "Чел 3"]
OUT = Path(__file__).resolve().parent.parent / "data" / "labels_clean.csv"
COLS = ["video_id", "raw_name", "violation", "timecode_sec", "split", "reviewer",
        "what_happens", "type_hypothesis", "how_visible"]


def norm(name) -> str:
    return str(name).strip().replace("X", "Х").replace("x", "Х")


def to_sec(s):
    if s is None or str(s).strip() in ("", "nan", "None"):
        return ""
    s = str(s)
    try:
        return str(float(s)).rstrip("0").rstrip(".")
    except ValueError:
        pass
    m = re.search(r"(\d+)\s*мин", s)
    sec = re.findall(r"(\d+)\s*сек", s)
    return str((int(m.group(1)) * 60 if m else 0) + (int(sec[0]) if sec else 0))


wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb["разметка"] if "разметка" in wb.sheetnames else wb.worksheets[0]
rows = [r for r in ws.iter_rows(values_only=True) if r and r[0]]
header = [str(c).strip() if c else "" for c in rows[0]]

if header[:2] == ["video_id", "raw_name"]:  # наш формат — просто пересохранить
    recs = [dict(zip(header, ["" if c is None else str(c) for c in r])) for r in rows[1:]]
    for r in recs:
        r["violation"] = str(int(float(r["violation"])))
        r["timecode_sec"] = to_sec(r["timecode_sec"])
else:  # исходная таблица организаторов
    recs = []
    for name, ts, label in (r[:3] for r in rows[1:]):
        recs.append({"video_id": norm(name), "raw_name": str(name).strip(),
                     "violation": "1" if str(label).strip().lower() == "нарушение" else "0",
                     "timecode_sec": to_sec(ts)})
    recs.sort(key=lambda r: int(re.search(r"\d+", r["video_id"]).group()))
    k = {"1": 0, "0": 1}
    for r in recs:
        if r["video_id"] in HOLDOUT:
            r["split"], r["reviewer"] = "holdout", "Чел 3 (в конце, только тип)"
        else:
            r["split"], r["reviewer"] = "train", PEOPLE[k[r["violation"]] % 3]
            k[r["violation"]] += 1

OUT.parent.mkdir(parents=True, exist_ok=True)
with OUT.open("w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore")
    w.writeheader()
    for r in recs:
        w.writerow({c: r.get(c, "") for c in COLS})

nv = sum(r["violation"] == "1" for r in recs)
print(f"OK: {OUT}  видео: {len(recs)}, с нарушением: {nv}, без: {len(recs) - nv}")
mine = [r["video_id"] for r in recs if r.get("reviewer") == "Чел 1"]
print(f"Чел 1 смотрит: {', '.join(mine)}")
